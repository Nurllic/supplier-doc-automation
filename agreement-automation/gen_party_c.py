
import os, sys, re
from datetime import date

SKILL_PATH = r"C:\Users\5532\.workbuddy\skills\word-doc-extractor\scripts"
sys.path.insert(0, SKILL_PATH)
if "extract_doc" in sys.modules:
    del sys.modules["extract_doc"]

QA_DIR = r".\agreement-automation"
sys.path.insert(0, QA_DIR)

from extract_doc import extract_table_by_signature, parse_quality_standard
from gen_agreements_new import parse_raw_material_new, parse_quality_standard_new, build_multi_quality_table, _extract_huohao_only
from generate_qa_agreement_v3 import (
    TEMPLATE_DOCX, get_platform_contact, replace_placeholder,
    rebuild_material_table, extract_spec, extract_part_number,
)
from docx import Document
import logging
logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

DATA_DIR = r".\data"
party_c_dir = os.path.join(DATA_DIR, "示例丙公司")

# Parse both raw material files
mats1 = parse_raw_material_new(os.path.join(party_c_dir, "materials/新增原辅料明细_v1.docx"))
mats2 = parse_raw_material_new(os.path.join(party_c_dir, "materials/新增原辅料明细表-平台A-张三-4.docx"))
all_materials = mats1 + mats2
for idx, m in enumerate(all_materials):
    m["序号"] = str(idx + 1).zfill(2)
logger.info(f"示例丙公司合并: {len(mats1)} + {len(mats2)} = {len(all_materials)} 物料")

# Parse all quality standards
std_dir = os.path.join(party_c_dir, "materials/质量标准_v1")
std_files = [os.path.join(std_dir, f) for f in os.listdir(std_dir) if f.endswith(".docx") and not f.startswith("~$")]

huohao_map = {}
for m in all_materials:
    code = m["物料编码"]
    cas_field = m.get("货号克隆号CAS", "")
    m_hh = re.search(r"货号[：:]\s*([^\n;]+)", cas_field)
    huohao_map[code] = _extract_huohao_only(cas_field) or None

def _match_std(std_path, mats):
    for m in mats:
        if m["物料编码"] in std_path:
            return m
    for m in mats:
        part_no = extract_part_number(m.get("货号克隆号CAS", ""))
        if part_no and part_no in std_path:
            return m
    for m in mats:
        mat_name = m.get("物料名称", "").lower()
        if mat_name and mat_name in os.path.basename(std_path).lower():
            return m
    return None

material_groups = []
for m in all_materials:
    code = m["物料编码"]
    huohao = huohao_map.get(code) or "/"
    name = m.get("物料中文名称") or m.get("物料名称") or ""
    spec = extract_spec(m.get("规格浓度", ""))
    material_groups.append({"code": huohao, "material_code": code, "name": name, "spec": spec, "indicators": []})

for sf in std_files:
    mat = _match_std(sf, all_materials)
    inds = parse_quality_standard_new(sf)
    if mat and inds:
        for g in material_groups:
            if g["material_code"] == mat["物料编码"]:
                g["indicators"] = inds
                break
        logger.info(f"  {os.path.basename(sf)} -> {mat['物料编码']} ({len(inds)} 指标)")
    elif not mat:
        logger.warning(f"  {os.path.basename(sf)} -> 未匹配到物料")

n_with = sum(1 for g in material_groups if g["indicators"])
logger.info(f"有质量标准: {n_with}/{len(material_groups)}")

# Generate document
doc = Document(TEMPLATE_DOCX)
supplier_name = "示例丙公司"
platform_name = "平台A"
oa_serial = "QA-20260715-002"

reviewer, phone, email = get_platform_contact(platform_name)
company_name = "示例甲公司"
company_address = "YY省YY市YY区YY路16号"
today = date.today().strftime("%Y年%m月%d日")

replace_placeholder(doc, "{1}", supplier_name)
replace_placeholder(doc, "{2}", reviewer)
replace_placeholder(doc, "{3}", phone)
replace_placeholder(doc, "{4}", email)
replace_placeholder(doc, "{5}", today)
replace_placeholder(doc, "{6}", oa_serial)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if "示例甲公司" in run.text:
                        run.text = run.text.replace("示例甲公司", company_name)
                    if "YY省YY市YY区YY路16号" in run.text:
                        run.text = run.text.replace("YY省YY市YY区YY路16号", company_address)

rebuild_material_table(doc, all_materials)
build_multi_quality_table(doc, material_groups, supplier_name, company_name)

out_path = os.path.join(party_c_dir, "agreements/v2_QA-20260715-002-【生物】-示例丙公司-10.docx")
doc.save(out_path)
logger.info(f"OK: {out_path}")
