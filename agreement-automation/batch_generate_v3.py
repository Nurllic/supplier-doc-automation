#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""批量生成质量保证协议 — 16物料"""
import os, sys, logging, json, re
from datetime import date
sys.path.insert(0, r'./agreement-automation')
from docx import Document
from docx.shared import Pt, Cm
from generate_qa_agreement_v3 import (
    TEMPLATE_DOCX, OUTPUT_DIR, TEST_RAW_MATERIAL,
    parse_raw_material, parse_quality_standard, extract_spec,
    replace_placeholder, get_platform_contact, sanitize_filename,
    _format_cell, _set_col_widths, _set_table_full_width, merge_cells_vertically,
    rebuild_material_table, build_multi_quality_table, get_party_a_info,
    build_output_filename, normalize_oa_serial, extract_part_number, fill_party_b_info,
)
QA_DIR = r'./email-database/QA'
logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

# 占位示例：真实物料编码/标准文件名已脱敏，实际使用时按业务数据维护
MATERIAL_STD_FILES = {
    "8800054594": "物料A原料企业质量标准.docx",
    "8800054595": "物料B原料企业质量标准.docx",
    "8800054603": "物料C原料企业质量标准.docx",
    "8800054942": "物料D原料企业质量标准.docx",
    "8800054596": "物料E原料企业质量标准.docx",
    "5101004417": "物料F原料企业质量标准.docx",
    "5800000595": "物料G原料企业质量标准.docx",
    "8800054600": "物料H原料企业质量标准.docx",
    "8800054606": "物料I原料企业质量标准.docx",
    "5101000220": "物料J原料企业质量标准.docx",
    "5101000219": "物料K原料企业质量标准.docx",
    "5800000616": "物料L原料企业质量标准.docx",
    "5101000215": "物料M原料企业质量标准.docx",
    "8800054601": "物料N原料企业质量标准.docx",
    "8800054952": "物料O原料企业质量标准.docx",
    "8800054602": "物料P原料企业质量标准.docx",
}


def batch_generate(raw_material_docx, supplier_name, platform_name="平台D",
                   oa_serial="QA-20260707-001", output_path=None,
                   material_filter=None, std_map=None, name_overrides=None,
                   party_a_contact=None, party_a_company=None,
                   master_overrides=None, party_b_info=None):
    doc = Document(TEMPLATE_DOCX)

    # 支持单个文件或文件列表；同一供应商可能来自多份原辅料明细
    raw_files = [raw_material_docx] if isinstance(raw_material_docx, str) else list(raw_material_docx)
    materials = []
    seen_codes = set()
    for rf in raw_files:
        for m in parse_raw_material(rf, material_filter=material_filter):
            code = m["物料编码"]
            if code in seen_codes:
                continue
            seen_codes.add(code)
            materials.append(m)

    # 主数据兜底：明细文件缺失（如纯标准升版）时，从物料主数据构建物料行
    if master_overrides:
        for code, entry in master_overrides.items():
            if code in seen_codes:
                continue
            if material_filter is not None and code not in material_filter:
                continue
            seen_codes.add(code)
            materials.append({
                "序号": str(len(materials) + 1).zfill(2),
                "物料编码": code,
                "物料名称": entry.get("name", code),
                "物料中文名称": entry.get("name", code),
                "规格浓度": entry.get("spec", ""),
                "主计量单位": entry.get("main_unit", ""),
                "货号克隆号CAS": entry.get("part_no", ""),
                "品牌产地": entry.get("brand", ""),
                "是否危化品": entry.get("dangerous", "否"),
                "物料属性": entry.get("property", ""),
                "储存条件": entry.get("storage", ""),
                "保质期": entry.get("shelf_life", ""),
                "开瓶次数": entry.get("open_times", "/"),
                "冻融次数": entry.get("freeze_thaw_times", "/"),
            })
    logger.info("共 %d 个物料", len(materials))

    # 从原辅料明细的货号/克隆号或CAS号列提取货号（支持多份明细合并、不同表头布局）
    huohao_map = {}
    for rf in raw_files:
        if rf.lower().endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(rf, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                hdr = [re.sub(r"\s+", "", str(c or "")) for c in rows[0]]
                code_col = next((i for i, h in enumerate(hdr) if "物料编码" in h), None)
                cas_col = next((i for i, h in enumerate(hdr) if "货号" in h), None)
                if code_col is not None and cas_col is not None:
                    for row in rows[1:]:
                        if row is None or len(row) <= max(code_col, cas_col):
                            continue
                        code = str(row[code_col] or "").strip()
                        if not code:
                            continue
                        part = extract_part_number(str(row[cas_col] or "").strip())
                        huohao_map[code] = part or None
            continue
        rm_doc = Document(rf)
        for t in rm_doc.tables:
            hdr = [re.sub(r"\s+", "", c.text) for c in t.rows[0].cells]
            code_col = next((i for i, h in enumerate(hdr) if "物料编码" in h), None)
            cas_col = next((i for i, h in enumerate(hdr) if "货号" in h), None)
            if code_col is None or cas_col is None:
                continue
            for ri in range(1, len(t.rows)):
                cells = [c.text.strip() for c in t.rows[ri].cells]
                if len(cells) <= max(code_col, cas_col) or not cells[code_col]:
                    continue
                code = cells[code_col]
                part = extract_part_number(cells[cas_col])
                huohao_map[code] = part or None  # None=无货号
            break

    # 质量标准映射：内置映射 + data/_party_c_std_map.json + 调用方传入的映射（后者优先）
    effective_std_map = dict(MATERIAL_STD_FILES)
    try:
        repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        with open(os.path.join(repo_root, "data", "_party_c_std_map.json"), encoding="utf-8") as _f:
            effective_std_map.update(json.load(_f))
    except Exception as e:
        logger.warning("读取 _party_c_std_map.json 失败: %s", e)
    if std_map:
        effective_std_map.update(std_map)

    material_groups = []
    for m in materials:
        code = m["物料编码"]
        huohao = huohao_map.get(code)
        if huohao is None and master_overrides and master_overrides.get(code, {}).get("part_no"):
            huohao = master_overrides[code]["part_no"]
        if huohao is None:
            huohao = "/"  # 无货号时显示/，不回退物料编码
        name = m["物料中文名称"] or m["物料名称"]
        if name_overrides and code in name_overrides and name_overrides[code]:
            # 覆盖值只用于表1"物料名称"列；表2"中/英文品名"始终优先物料中文名称
            m["物料名称"] = name_overrides[code]
            if not m.get("物料中文名称") or m["物料中文名称"] in ("", "/"):
                m["物料中文名称"] = name_overrides[code]
            name = m["物料中文名称"] or name_overrides[code]
        spec = extract_spec(m["规格浓度"])
        inds = []
        std_file = effective_std_map.get(code)
        if std_file:
            sp = std_file if os.path.isabs(std_file) else os.path.join(QA_DIR, std_file)
            if os.path.exists(sp) and os.path.getsize(sp) > 0:
                try:
                    inds = parse_quality_standard(sp)
                    if isinstance(inds, dict):
                        inds = [{"检验项目": k, "标准值": v.get("标准值", "/"), "测试方法": v.get("测试方法", "/")}
                                for k, v in inds.items()]
                    logger.info("  ok %s: %d 指标", name, len(inds))
                except Exception as e:
                    logger.warning("  fail %s: %s", name, e)
            else:
                sz = os.path.getsize(sp) if os.path.exists(sp) else 0
                logger.warning("  skip %s: %d B", name, sz)
        material_groups.append({"code": huohao, "material_code": code, "name": name, "spec": spec, "indicators": inds})
    n_with = sum(1 for g in material_groups if g["indicators"])
    n_wo = sum(1 for g in material_groups if not g["indicators"])
    logger.info("有标准 %d / 无标准 %d", n_with, n_wo)

    # 甲方公司/联系人：钉钉AI表格确认的“甲方公司/甲方联系人”优先，否则按平台默认
    company_name, company_address, reviewer, phone, email = get_party_a_info(
        platform_name, party_a_contact, party_a_company
    )
    logger.info("甲方: %s | 审核人: %s", company_name, reviewer)
    today = date.today().strftime("%Y年%m月%d日")
    replace_placeholder(doc, "{1}", supplier_name)
    replace_placeholder(doc, "{2}", reviewer)
    replace_placeholder(doc, "{3}", phone)
    replace_placeholder(doc, "{4}", email)
    replace_placeholder(doc, "{5}", today)
    replace_placeholder(doc, "{6}", normalize_oa_serial(oa_serial))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if "示例甲公司" in run.text:
                            run.text = run.text.replace("示例甲公司", company_name)
                        if "YY省YY市YY区YY路16号" in run.text:
                            run.text = run.text.replace("YY省YY市YY区YY路16号", company_address)
    fill_party_b_info(doc, party_b_info)
    rebuild_material_table(doc, materials)
    build_multi_quality_table(doc, material_groups, supplier_name, company_name)
    if output_path is None:
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        tag = "平台C" if ("示例丙公司" in (party_a_company or "")) or (not party_a_company and platform_name == "平台C") else "生物"
        output_path = os.path.join(OUTPUT_DIR, build_output_filename(oa_serial, tag, supplier_name))
    doc.save(output_path)
    logger.info("OK: %s", output_path)
    return output_path

if __name__ == "__main__":
    rf = os.path.join(QA_DIR, "新增原辅料明细.docx")
    if not os.path.exists(rf):
        rf = TEST_RAW_MATERIAL
    out = batch_generate(
        raw_material_docx=rf,
        supplier_name="供应商A公司",
        platform_name="平台D",
        oa_serial="QA-20260707-001",
    )
    print("\n输出:", out)
