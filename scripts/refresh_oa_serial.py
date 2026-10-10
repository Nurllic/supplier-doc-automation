#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""OA流水号回填刷新：把已生成协议的页眉流水号、文件名、history 更新为正式OA流水号。
用法: python refresh_oa_serial.py --supplier 供应商B公司 --serial QA-DOC-20260812-001 [--dry-run]
"""

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "agreement-automation"))

from docx import Document  # noqa: E402

from generate_qa_agreement_v3 import build_output_filename, normalize_oa_serial  # noqa: E402

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(REPO_ROOT, "data")


def _find_supplier_dir(supplier):
    """新结构：data/{甲方公司}/{供应商}"""
    for company in ("示例丙公司", "示例甲公司"):
        p = os.path.join(DATA_DIR, company, supplier)
        if os.path.exists(os.path.join(p, "history.json")):
            return p
    return os.path.join(DATA_DIR, supplier)


def refresh(supplier, oa_serial, dry_run=False):
    supplier_dir = _find_supplier_dir(supplier)
    hp = os.path.join(supplier_dir, "history.json")
    if not os.path.exists(hp):
        return False, "未找到该供应商的 history.json"
    with open(hp, encoding="utf-8") as f:
        history = json.load(f)
    versions = history.get("versions", [])
    if not versions:
        return False, "history 无版本记录"
    version = versions[-1]
    old_serial = normalize_oa_serial(version.get("oa_serial") or version.get("oa_serial_temp") or "")
    new_serial = normalize_oa_serial(oa_serial)
    if old_serial == new_serial:
        return True, "流水号未变化，无需刷新"
    old_tokens = {old_serial}
    raw_serial = str(version.get("oa_serial") or version.get("oa_serial_temp") or "")
    if raw_serial:
        old_tokens.add(raw_serial)
    if raw_serial.startswith("QA-"):
        old_tokens.add(raw_serial[3:])

    def replace_serial(text):
        for tok in old_tokens:
            if tok and tok in text:
                text = text.replace(tok, new_serial)
        return text
    out_rel = version.get("output_file", "")
    old_path = os.path.join(supplier_dir, out_rel)
    if not os.path.exists(old_path):
        return False, f"协议文件不存在: {old_path}"
    tag = "平台C" if ("平台C" in out_rel or version.get("platform") == "平台C") else "生物"
    new_name = build_output_filename(new_serial, tag, supplier)
    new_path = os.path.join(os.path.dirname(old_path), new_name)
    if dry_run:
        return True, f"将重命名: {os.path.basename(old_path)} -> {new_name}"

    doc = Document(old_path)
    for section in doc.sections:
        for para in section.header.paragraphs:
            for run in para.runs:
                run.text = replace_serial(run.text)
    for para in doc.paragraphs:
        for run in para.runs:
            run.text = replace_serial(run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.text = replace_serial(run.text)
    doc.save(new_path)
    os.remove(old_path)

    version["oa_serial"] = new_serial
    version["output_file"] = os.path.relpath(new_path, supplier_dir)
    with open(hp, "w", encoding="utf-8") as f:
        json.dump(history, f, ensure_ascii=False, indent=2)
    return True, f"已刷新: {new_name}"


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--supplier", required=True)
    parser.add_argument("--serial", required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    ok, msg = refresh(args.supplier, args.serial, dry_run=args.dry_run)
    print(("OK: " if ok else "FAIL: ") + msg)
    sys.exit(0 if ok else 1)
